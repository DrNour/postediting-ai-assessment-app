import io
import json
import zipfile
import difflib
import html
import math
import secrets as pysecrets
import string
from collections import Counter

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from docx import Document
from docx.shared import RGBColor
from supabase import create_client
from metrics import compare_postedit_with_raw_mt, build_research_metrics_payload
from modules.auth import require_teacher_access
from modules.task_mode import (
    POST_EDITING,
    POST_EDITING_ONLY_METRIC_FIELDS,
    TRANSLATION,
    TASK_OPTIONS,
    is_translation,
    make_translation_metrics_task_appropriate,
    normalize_task_type,
    student_output_label,
    task_instruction,
    task_type_label,
)


def install_student_paste_guard():
    """Apply browser-side integrity controls to assessed response boxes.

    The guard blocks paste/drop and common dictation-style insertion events. On
    phones and tablets the assessed response boxes are made read-only so students
    must complete the task on a desktop/laptop. This remains a deterrent rather
    than a mathematically foolproof proctoring mechanism: operating-system tools
    can sometimes make dictated text look like ordinary keyboard input.
    """
    st.info(
        "Academic integrity mode is active: paste, drag/drop, and detected voice "
        "dictation are blocked. Assessed responses must be completed on a "
        "desktop/laptop and typed directly into the app."
    )

    guard_js = r"""
    <script>
    (() => {
      let rootDoc;
      let rootWin;
      try {
        rootDoc = window.parent.document;
        rootWin = window.parent;
      } catch (e) {
        rootDoc = document;
        rootWin = window;
      }

      const protectedLabels = new Set(["Translation box", "Post-editing box"]);
      const mobileRe = /Android|webOS|iPhone|iPad|iPod|BlackBerry|IEMobile|Opera Mini|Mobile|Tablet/i;
      const isMobile = mobileRe.test((rootWin.navigator && rootWin.navigator.userAgent) || "");

      const isProtected = (el) =>
        el && el.tagName === "TEXTAREA" && protectedLabels.has(el.getAttribute("aria-label"));

      const addMessage = (el, text, kind) => {
        const host = el.closest('[data-testid="stTextArea"]') || el.parentElement;
        if (!host) return;
        const cls = "eduapp-integrity-" + kind;
        if (host.querySelector("." + cls)) return;
        const msg = rootDoc.createElement("div");
        msg.className = cls;
        msg.textContent = text;
        msg.style.fontSize = "0.9rem";
        msg.style.fontWeight = "600";
        msg.style.marginTop = "0.35rem";
        msg.style.padding = "0.45rem 0.6rem";
        msg.style.borderRadius = "0.35rem";
        msg.style.background = "rgba(255, 193, 7, 0.14)";
        msg.style.border = "1px solid rgba(255, 193, 7, 0.45)";
        host.appendChild(msg);
      };

      const flashBlocked = (el, reason) => {
        const oldTitle = el.getAttribute("title") || "";
        el.setAttribute("title", reason);
        el.style.outline = "2px solid var(--st-primary-color, #ff4b4b)";
        addMessage(el, reason, "blocked");
        rootWin.setTimeout(() => {
          el.style.outline = "";
          if (oldTitle) el.setAttribute("title", oldTitle);
          else el.removeAttribute("title");
        }, 1600);
      };

      const block = (event, reason) => {
        event.preventDefault();
        event.stopPropagation();
        if (event.stopImmediatePropagation) event.stopImmediatePropagation();
        flashBlocked(event.currentTarget || event.target, reason);
        return false;
      };

      const protect = (el) => {
        if (!isProtected(el) || el.dataset.eduappIntegrityGuard === "1") return;
        el.dataset.eduappIntegrityGuard = "1";
        el.setAttribute("autocomplete", "off");
        el.setAttribute("autocapitalize", "off");
        el.setAttribute("spellcheck", "false");

        // Strongest practical protection against phone-keyboard dictation:
        // assessed response entry is disabled on phones/tablets.
        if (isMobile) {
          el.readOnly = true;
          el.setAttribute("inputmode", "none");
          el.setAttribute("placeholder", "Use a desktop or laptop for this assessed task.");
          addMessage(
            el,
            "Phone/tablet input is disabled for this assessed task. Please use a desktop or laptop.",
            "mobile"
          );
        }

        el.addEventListener("paste", (event) =>
          block(event, "Pasting is disabled. Type the response directly."), true);
        el.addEventListener("drop", (event) =>
          block(event, "Dropping external text is disabled."), true);

        el.addEventListener("beforeinput", (event) => {
          const t = event.inputType || "";
          if (t === "insertFromPaste" || t === "insertFromDrop") {
            block(event, "Pasting or dropping external text is disabled.");
            return;
          }
          if (t === "insertFromDictation") {
            block(event, "Voice dictation is disabled for this assessed task.");
            return;
          }

          // Many mobile/OS dictation systems expose a whole phrase as one trusted
          // insertText/replacement event. Block unusually large single-event inserts
          // while leaving normal typing and IME composition alone.
          const data = typeof event.data === "string" ? event.data : "";
          const largeChunk = !event.isComposing && data.length >= 8;
          if ((t === "insertText" || t === "insertReplacementText") && largeChunk) {
            block(event, "Large one-step text insertion/voice dictation is disabled. Type normally.");
          }
        }, true);

        el.addEventListener("keydown", (event) => {
          const key = (event.key || "").toLowerCase();
          const pasteShortcut = (event.ctrlKey || event.metaKey) && key === "v";
          const shiftInsert = event.shiftKey && event.key === "Insert";
          if (pasteShortcut || shiftInsert) {
            block(event, "Pasting is disabled. Type the response directly.");
          }
        }, true);
      };

      const scan = () => rootDoc.querySelectorAll("textarea").forEach(protect);
      scan();

      if (!rootWin.__eduappIntegrityGuardObserver) {
        const observer = new rootWin.MutationObserver(scan);
        observer.observe(rootDoc.documentElement, { childList: true, subtree: true });
        rootWin.__eduappIntegrityGuardObserver = observer;
      }
    })();
    </script>
    """

    # components.html executes JavaScript in a small iframe; the script then
    # targets the parent Streamlit document. Height zero keeps it invisible.
    components.html(guard_js, height=0, width=0)


# ============================================================
# Supabase connection
# ============================================================

@st.cache_resource
def get_supabase_client():
    try:
        url = st.secrets["SUPABASE_URL"]
        key = st.secrets["SUPABASE_KEY"]
        return create_client(url, key)
    except Exception:
        st.error(
            "Supabase is not configured. Add SUPABASE_URL and SUPABASE_KEY "
            "to Streamlit Secrets."
        )
        st.stop()



# ============================================================
# Supabase storage functions
# ============================================================

def load_assignments():
    try:
        response = (
            get_supabase_client().table("assignments")
            .select("*")
            .order("created_at", desc=True)
            .execute()
        )

        return pd.DataFrame(response.data or [])

    except Exception as error:
        st.error("Could not read the assignments table from Supabase.")
        st.info(
            "Check that the table named 'assignments' exists in the public schema "
            "and that Row Level Security is disabled or policies allow SELECT access."
        )
        st.code(str(error))
        return pd.DataFrame()


def save_assignment(assignment):
    try:
        return get_supabase_client().table("assignments").insert(assignment).execute()

    except Exception as error:
        st.error("Could not save the assignment to Supabase.")
        st.write("The assignment data being sent was:")
        st.json(assignment)
        st.write("Supabase error:")
        st.code(str(error))
        st.stop()


def update_assignment(assignment_id, updates):
    try:
        return (
            get_supabase_client().table("assignments")
            .update(updates)
            .eq("assignment_id", safe_text(assignment_id))
            .execute()
        )
    except Exception as error:
        st.error("Could not update the assignment in Supabase.")
        st.code(str(error))
        return None


def delete_assignment(assignment_id):
    """Delete an assignment while preserving already-submitted research records."""
    try:
        # Drafts are unfinished work tied to the exercise and should disappear with it.
        try:
            get_supabase_client().table("student_drafts").delete().eq(
                "assignment_id", safe_text(assignment_id)
            ).execute()
        except Exception:
            pass

        return (
            get_supabase_client().table("assignments")
            .delete()
            .eq("assignment_id", safe_text(assignment_id))
            .execute()
        )
    except Exception as error:
        st.error("Could not delete the assignment from Supabase.")
        st.code(str(error))
        return None


def parse_audience_rules(value):
    """Return normalized [{group, code}] audience rules from JSON/list values."""
    if value is None or value == "":
        return []
    data = value
    if isinstance(value, str):
        try:
            data = json.loads(value)
        except Exception:
            return []
    if not isinstance(data, list):
        return []
    rules = []
    for item in data:
        if not isinstance(item, dict):
            continue
        group = safe_text(item.get("group"))
        code = normalize_access_code(item.get("code"))
        if group and code:
            rules.append({"group": group, "code": code})
    return rules





def normalize_access_code(value):
    """Normalize access codes as text so numeric codes like 123 stay '123'."""
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().upper().replace(" ", "")


def generate_access_code(length=6):
    """Generate a short classroom access code that avoids ambiguous characters."""
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    return "".join(pysecrets.choice(alphabet) for _ in range(length))


def audience_rules_from_editor(editor_value):
    """Normalize the class/group table into [{group, code}] rules.

    Blank access codes are generated automatically. Completely blank rows are ignored.
    """
    if editor_value is None:
        return [], []
    try:
        df = pd.DataFrame(editor_value).copy()
    except Exception:
        return [], ["Could not read the class/group table."]

    if df.empty:
        return [], []

    # Accept both current UI column labels and any older internal names.
    rename_map = {}
    for col in df.columns:
        key = safe_text(col).strip().casefold()
        if key in {"class / group", "class/group", "group", "class", "group name"}:
            rename_map[col] = "group"
        elif key in {"access code", "code"}:
            rename_map[col] = "code"
    df = df.rename(columns=rename_map)
    if "group" not in df.columns:
        df["group"] = ""
    if "code" not in df.columns:
        df["code"] = ""

    rules = []
    errors = []
    seen_groups = set()
    seen_codes = set()

    for row_no, row in enumerate(df.to_dict("records"), start=1):
        group = safe_text(row.get("group")).strip()
        code = normalize_access_code(row.get("code"))
        if not group and not code:
            continue
        if not group:
            errors.append(f"Row {row_no}: enter a class/group name, or delete the row.")
            continue
        if not code:
            code = generate_access_code()

        group_key = group.casefold()
        code_key = code.casefold()
        if group_key in seen_groups:
            errors.append(f"Row {row_no}: class/group '{group}' is duplicated.")
            continue
        if code_key in seen_codes:
            errors.append(f"Row {row_no}: access code '{code}' is duplicated.")
            continue

        seen_groups.add(group_key)
        seen_codes.add(code_key)
        rules.append({"group": group, "code": code})

    return rules, errors


def load_submissions():
    try:
        response = (
            get_supabase_client().table("submissions")
            .select("*")
            .order("submitted_at", desc=True)
            .execute()
        )

        return pd.DataFrame(response.data or [])

    except Exception as error:
        st.error("Could not read the submissions table from Supabase.")
        st.info(
            "Check that the table named 'submissions' exists in the public schema "
            "and that Row Level Security is disabled or policies allow SELECT access."
        )
        st.code(str(error))
        return pd.DataFrame()

def clean_value_for_supabase(value):
    """
    Converts values that Supabase/PostgREST may reject.
    """
    try:
        if value is None:
            return None

        if isinstance(value, float):
            if math.isnan(value) or math.isinf(value):
                return None
            return value

        # Handles pandas/numpy missing values
        if pd.isna(value):
            return None

        return value

    except Exception:
        return value


def save_submission(submission):
    """
    Saves a student submission and shows the real Supabase error if insertion fails.
    """

    clean_submission = {
        key: clean_value_for_supabase(value)
        for key, value in submission.items()
    }

    try:
        return get_supabase_client().table("submissions").insert(clean_submission).execute()

    except Exception as error:
        st.error("Could not save the submission to Supabase.")

        error_text = str(error)
        if "task_type" in error_text.lower():
            st.warning(
                "Your submissions table does not yet have the task_type column. "
                "Run supabase/migrations/001_add_task_type.sql once in the Supabase SQL Editor."
            )

        st.write("These are the columns the app is trying to send:")
        st.json(sorted(list(clean_submission.keys())))

        st.write("This is the full Supabase error:")
        st.code(error_text)

        st.stop()



def load_student_draft(assignment_id, task_type, student_id):
    """Return the most recent saved draft for one student/task, if present."""
    sid = safe_text(student_id)
    if not sid:
        return None
    try:
        response = (
            get_supabase_client().table("student_drafts")
            .select("*")
            .eq("assignment_id", safe_text(assignment_id))
            .eq("task_type", normalize_task_type(task_type))
            .eq("student_id", sid)
            .order("updated_at", desc=True)
            .limit(1)
            .execute()
        )
        rows = response.data or []
        return rows[0] if rows else None
    except Exception as error:
        st.error("Could not load the saved draft from Supabase.")
        if "student_drafts" in str(error).lower():
            st.warning(
                "Run the student_drafts migration in the Supabase SQL Editor once, then try again."
            )
        st.code(str(error))
        return None


def save_student_draft(draft):
    """Save or replace one student's draft without creating a final submission."""
    clean_draft = {
        key: clean_value_for_supabase(value)
        for key, value in draft.items()
    }
    client = get_supabase_client()
    try:
        # Delete+insert avoids depending on a particular supabase-py upsert signature.
        (
            client.table("student_drafts")
            .delete()
            .eq("assignment_id", clean_draft["assignment_id"])
            .eq("task_type", clean_draft["task_type"])
            .eq("student_id", clean_draft["student_id"])
            .execute()
        )
        return client.table("student_drafts").insert(clean_draft).execute()
    except Exception as error:
        st.error("Could not save the draft to Supabase.")
        if "student_drafts" in str(error).lower():
            st.warning(
                "Run the student_drafts migration in the Supabase SQL Editor once, then try again."
            )
        st.code(str(error))
        return None


def delete_student_draft(assignment_id, task_type, student_id):
    """Remove a draft after the final submission has been stored successfully."""
    sid = safe_text(student_id)
    if not sid:
        return
    try:
        (
            get_supabase_client().table("student_drafts")
            .delete()
            .eq("assignment_id", safe_text(assignment_id))
            .eq("task_type", normalize_task_type(task_type))
            .eq("student_id", sid)
            .execute()
        )
    except Exception:
        # A stale draft is not serious enough to invalidate a successful submission.
        pass

def update_submission_review(submission_id, teacher_score, teacher_feedback):
    return (
        get_supabase_client().table("submissions")
        .update(
            {
                "teacher_score": teacher_score,
                "teacher_feedback": teacher_feedback,
            }
        )
        .eq("submission_id", submission_id)
        .execute()
    )


# ============================================================
# Text and metric helpers
# ============================================================

def safe_text(text):
    if text is None:
        return ""

    if isinstance(text, float) and math.isnan(text):
        return ""

    return str(text).strip()


def word_count(text):
    text = safe_text(text)

    if not text:
        return 0

    return len(text.split())


def lexical_cosine_similarity(text_a, text_b):
    words_a = safe_text(text_a).lower().split()
    words_b = safe_text(text_b).lower().split()

    if not words_a or not words_b:
        return None

    counter_a = Counter(words_a)
    counter_b = Counter(words_b)

    common_words = set(counter_a.keys()) & set(counter_b.keys())

    numerator = sum(counter_a[word] * counter_b[word] for word in common_words)

    sum_a = sum(value ** 2 for value in counter_a.values())
    sum_b = sum(value ** 2 for value in counter_b.values())

    denominator = math.sqrt(sum_a) * math.sqrt(sum_b)

    if denominator == 0:
        return None

    return round(numerator / denominator, 4)


@st.cache_resource
def load_sentence_transformer_model():
    try:
        from sentence_transformers import SentenceTransformer

        return SentenceTransformer(
            "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
        )
    except Exception:
        return None


def semantic_cosine_similarity(text_a, text_b):
    text_a = safe_text(text_a)
    text_b = safe_text(text_b)

    if not text_a or not text_b:
        return None

    model = load_sentence_transformer_model()

    if model is None:
        return lexical_cosine_similarity(text_a, text_b)

    embeddings = model.encode(
        [text_a, text_b],
        normalize_embeddings=True,
    )

    score = float(embeddings[0] @ embeddings[1])

    return round(score, 4)


def cosine_similarity(text_a, text_b, use_semantic=False):
    if use_semantic:
        return semantic_cosine_similarity(text_a, text_b)

    return lexical_cosine_similarity(text_a, text_b)


def edit_distance_ratio(text_a, text_b):
    text_a = safe_text(text_a)
    text_b = safe_text(text_b)

    if not text_a and not text_b:
        return 0.0

    ratio = difflib.SequenceMatcher(None, text_a, text_b).ratio()

    return round(1 - ratio, 4)


def length_ratio(candidate, reference):
    candidate_words = word_count(candidate)
    reference_words = word_count(reference)

    if reference_words == 0:
        return None

    return round(candidate_words / reference_words, 3)


def reference_based_scores(candidate, reference):
    candidate = safe_text(candidate)
    reference = safe_text(reference)

    if not candidate or not reference:
        return {
            "bleu": None,
            "chrf": None,
            "ter": None,
        }

    try:
        from sacrebleu.metrics import BLEU, CHRF, TER

        bleu = BLEU()
        chrf = CHRF()
        ter = TER()

        return {
            "bleu": round(bleu.sentence_score(candidate, [reference]).score, 3),
            "chrf": round(chrf.sentence_score(candidate, [reference]).score, 3),
            "ter": round(ter.sentence_score(candidate, [reference]).score, 3),
        }

    except Exception:
        return {
            "bleu": None,
            "chrf": None,
            "ter": None,
        }


def compute_bert_score(candidate, reference, language="en"):
    candidate = safe_text(candidate)
    reference = safe_text(reference)

    if not candidate or not reference:
        return None

    try:
        from bert_score import score

        _, _, f1 = score(
            [candidate],
            [reference],
            lang=language,
            verbose=False,
            rescale_with_baseline=False,
        )

        return round(float(f1[0]), 4)

    except Exception:
        return None


def make_track_changes_html(original_text, edited_text):
    original_words = safe_text(original_text).split()
    edited_words = safe_text(edited_text).split()

    matcher = difflib.SequenceMatcher(None, original_words, edited_words)

    output = []

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for word in original_words[i1:i2]:
                output.append(
                    f'<span class="same-word">{html.escape(word)}</span>'
                )

        elif tag == "delete":
            for word in original_words[i1:i2]:
                output.append(
                    f'<span class="deleted-word">{html.escape(word)}</span>'
                )

        elif tag == "insert":
            for word in edited_words[j1:j2]:
                output.append(
                    f'<span class="added-word">{html.escape(word)}</span>'
                )

        elif tag == "replace":
            for word in original_words[i1:i2]:
                output.append(
                    f'<span class="deleted-word">{html.escape(word)}</span>'
                )

            for word in edited_words[j1:j2]:
                output.append(
                    f'<span class="added-word">{html.escape(word)}</span>'
                )

    return " ".join(output)


def calculate_edit_summary(original_text, edited_text):
    original_words = safe_text(original_text).split()
    edited_words = safe_text(edited_text).split()

    matcher = difflib.SequenceMatcher(None, original_words, edited_words)

    inserted = 0
    deleted = 0
    replaced = 0
    unchanged = 0

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            unchanged += i2 - i1
        elif tag == "delete":
            deleted += i2 - i1
        elif tag == "insert":
            inserted += j2 - j1
        elif tag == "replace":
            deleted += i2 - i1
            inserted += j2 - j1
            replaced += max(i2 - i1, j2 - j1)

    return {
        "inserted_words": inserted,
        "deleted_words": deleted,
        "replaced_segments": replaced,
        "unchanged_words": unchanged,
    }


def compare_mt_and_postedit(
    raw_mt,
    post_edited_text,
    use_semantic_cosine=False,
    use_bert=False,
    bert_language="en",
):
    scores = reference_based_scores(post_edited_text, raw_mt)

    cosine_method = (
        "semantic_sentence_transformer"
        if use_semantic_cosine
        else "fast_lexical"
    )

    return {
        "mt_pe_cosine_similarity": cosine_similarity(
            raw_mt,
            post_edited_text,
            use_semantic=use_semantic_cosine,
        ),
        "mt_pe_cosine_method": cosine_method,
        "mt_pe_edit_distance_ratio": edit_distance_ratio(
            raw_mt,
            post_edited_text,
        ),
        "mt_pe_length_ratio": length_ratio(post_edited_text, raw_mt),
        "mt_pe_bleu": scores["bleu"],
        "mt_pe_chrf": scores["chrf"],
        "mt_pe_ter": scores["ter"],
        "mt_pe_bertscore_f1": compute_bert_score(
            post_edited_text,
            raw_mt,
            language=bert_language,
        ) if use_bert else None,
    }


def compare_postedit_and_reference(
    post_edited_text,
    reference_translation,
    use_semantic_cosine=False,
    use_bert=False,
    bert_language="en",
):
    reference_translation = safe_text(reference_translation)

    if not reference_translation:
        return {
            "pe_reference_cosine_similarity": None,
            "pe_reference_cosine_method": "",
            "pe_reference_length_ratio": None,
            "pe_reference_bleu": None,
            "pe_reference_chrf": None,
            "pe_reference_ter": None,
            "pe_reference_bertscore_f1": None,
        }

    scores = reference_based_scores(post_edited_text, reference_translation)

    cosine_method = (
        "semantic_sentence_transformer"
        if use_semantic_cosine
        else "fast_lexical"
    )

    return {
        "pe_reference_cosine_similarity": cosine_similarity(
            post_edited_text,
            reference_translation,
            use_semantic=use_semantic_cosine,
        ),
        "pe_reference_cosine_method": cosine_method,
        "pe_reference_length_ratio": length_ratio(
            post_edited_text,
            reference_translation,
        ),
        "pe_reference_bleu": scores["bleu"],
        "pe_reference_chrf": scores["chrf"],
        "pe_reference_ter": scores["ter"],
        "pe_reference_bertscore_f1": compute_bert_score(
            post_edited_text,
            reference_translation,
            language=bert_language,
        ) if use_bert else None,
    }


def build_quality_warnings(
    mt_pe_metrics,
    reference_metrics,
    output_word_count,
    task_type=POST_EDITING,
):
    warnings = []

    if output_word_count < 5:
        noun = "Student translation" if is_translation(task_type) else "Post-edited text"
        warnings.append(f"{noun} is very short.")

    if not is_translation(task_type):
        mt_pe_cosine = mt_pe_metrics.get("mt_pe_cosine_similarity")
        edit_ratio = mt_pe_metrics.get("mt_pe_edit_distance_ratio")

        if mt_pe_cosine is not None and mt_pe_cosine >= 0.95:
            warnings.append("Post-edited text is extremely close to the raw MT.")

        if edit_ratio is not None and edit_ratio < 0.05:
            warnings.append("Very little editing detected.")

    reference_cosine = reference_metrics.get("pe_reference_cosine_similarity")

    if reference_cosine is not None and reference_cosine < 0.50:
        warnings.append("Low similarity with the reference translation.")

    if not warnings:
        return "No automatic warnings."

    return " | ".join(warnings)


def metrics_to_dataframe(metrics):
    return pd.DataFrame(
        [{"Metric": key, "Value": value} for key, value in metrics.items()]
    )


# ============================================================
# Word export helpers
# ============================================================

def clean_filename(text):
    text = safe_text(text) or "submission"

    for char in '<>:"/\\|?*':
        text = text.replace(char, "_")

    return text.replace(" ", "_")[:80]


def add_docx_section(document, title, text):
    document.add_heading(title, level=2)
    document.add_paragraph(safe_text(text))


def add_track_changes_to_docx(document, original_text, edited_text):
    document.add_heading("Track Changes Style Preview", level=2)

    paragraph = document.add_paragraph()

    original_words = safe_text(original_text).split()
    edited_words = safe_text(edited_text).split()

    matcher = difflib.SequenceMatcher(None, original_words, edited_words)

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for word in original_words[i1:i2]:
                run = paragraph.add_run(word + " ")
                run.font.color.rgb = RGBColor(17, 24, 39)

        elif tag == "delete":
            for word in original_words[i1:i2]:
                run = paragraph.add_run(word + " ")
                run.font.strike = True
                run.font.color.rgb = RGBColor(153, 27, 27)

        elif tag == "insert":
            for word in edited_words[j1:j2]:
                run = paragraph.add_run(word + " ")
                run.bold = True
                run.font.color.rgb = RGBColor(6, 95, 70)

        elif tag == "replace":
            for word in original_words[i1:i2]:
                run = paragraph.add_run(word + " ")
                run.font.strike = True
                run.font.color.rgb = RGBColor(153, 27, 27)

            for word in edited_words[j1:j2]:
                run = paragraph.add_run(word + " ")
                run.bold = True
                run.font.color.rgb = RGBColor(6, 95, 70)


def create_submission_docx(submission):
    document = Document()
    task_type = normalize_task_type(submission.get("task_type"))
    output_heading = student_output_label(task_type)

    document.add_heading(f"Student {task_type_label(task_type)} Submission", level=1)

    document.add_heading("Student Information", level=2)
    document.add_paragraph(f"Student ID: {safe_text(submission.get('student_id'))}")
    document.add_paragraph(f"Student name: {safe_text(submission.get('student_name'))}")
    document.add_paragraph(f"Submitted at: {safe_text(submission.get('submitted_at'))}")

    document.add_heading("Assignment Information", level=2)
    document.add_paragraph(
        f"Assignment: {safe_text(submission.get('assignment_title'))}"
    )
    document.add_paragraph(f"Task type: {task_type_label(task_type)}")

    add_docx_section(document, "Source Text", submission.get("source_text"))

    raw_mt = safe_text(submission.get("machine_translation"))
    if raw_mt:
        mt_heading = (
            "Raw Machine Translation"
            if not is_translation(task_type)
            else "Raw Machine Translation (teacher/research record; hidden from student)"
        )
        add_docx_section(document, mt_heading, raw_mt)

    add_docx_section(
        document,
        "Reference Translation",
        submission.get("reference_translation"),
    )
    add_docx_section(
        document,
        output_heading,
        submission.get("post_edited_text"),
    )

    if not is_translation(task_type):
        add_track_changes_to_docx(
            document,
            submission.get("machine_translation"),
            submission.get("post_edited_text"),
        )

    document.add_heading("Automatic Metrics", level=2)

    common_metrics = [
        ("Source word count", submission.get("source_word_count")),
        ("Student-output word count", submission.get("pe_word_count")),
        (
            "Output-reference cosine similarity",
            submission.get("pe_reference_cosine_similarity"),
        ),
        ("Output-reference cosine method", submission.get("pe_reference_cosine_method")),
        ("Output-reference length ratio", submission.get("pe_reference_length_ratio")),
        ("Output-reference BLEU", submission.get("pe_reference_bleu")),
        ("Output-reference chrF", submission.get("pe_reference_chrf")),
        ("Output-reference TER", submission.get("pe_reference_ter")),
        ("Output-reference BERTScore F1", submission.get("pe_reference_bertscore_f1")),
        ("Research mode", submission.get("research_mode")),
        ("Advanced metrics status", submission.get("advanced_metrics_status")),
        ("Raw MT quality BLEU", submission.get("raw_mt_quality_bleu")),
        ("Raw MT quality chrF", submission.get("raw_mt_quality_chrf")),
        ("Raw MT quality TER", submission.get("raw_mt_quality_ter")),
        ("Student-output quality BLEU", submission.get("pe_quality_bleu")),
        ("Student-output quality chrF", submission.get("pe_quality_chrf")),
        ("Student-output quality TER", submission.get("pe_quality_ter")),
        ("Human-translation quality BLEU", submission.get("ht_quality_bleu")),
        ("Human-translation quality chrF", submission.get("ht_quality_chrf")),
        ("Human-translation quality TER", submission.get("ht_quality_ter")),
        ("Raw MT quality COMET", submission.get("raw_mt_quality_comet")),
        ("Student-output quality COMET", submission.get("pe_quality_comet")),
        ("Human-translation quality COMET", submission.get("ht_quality_comet")),
        ("Automatic interpretation", submission.get("mt_pe_interpretation")),
        ("Quality warnings", submission.get("quality_warnings")),
    ]

    post_editing_metrics = [
        ("Inserted words", submission.get("inserted_words")),
        ("Deleted words", submission.get("deleted_words")),
        ("Replaced segments", submission.get("replaced_segments")),
        ("Unchanged words", submission.get("unchanged_words")),
        ("MT word count", submission.get("mt_word_count")),
        ("MT-PE cosine similarity", submission.get("mt_pe_cosine_similarity")),
        ("MT-PE cosine method", submission.get("mt_pe_cosine_method")),
        ("MT-PE edit-distance ratio", submission.get("mt_pe_edit_distance_ratio")),
        ("MT-PE length ratio", submission.get("mt_pe_length_ratio")),
        ("MT-PE BLEU", submission.get("mt_pe_bleu")),
        ("MT-PE chrF", submission.get("mt_pe_chrf")),
        ("MT-PE TER", submission.get("mt_pe_ter")),
        ("MT-PE BERTScore F1", submission.get("mt_pe_bertscore_f1")),
    ]

    metrics = common_metrics
    if not is_translation(task_type):
        metrics = post_editing_metrics + common_metrics

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    header_cells[0].text = "Metric"
    header_cells[1].text = "Value"

    for metric, value in metrics:
        row_cells = table.add_row().cells
        row_cells[0].text = safe_text(metric)
        row_cells[1].text = safe_text(value)

    document.add_heading("Teacher Review", level=2)
    document.add_paragraph(
        f"Teacher score: {safe_text(submission.get('teacher_score'))}"
    )
    document.add_paragraph(
        f"Teacher feedback: {safe_text(submission.get('teacher_feedback'))}"
    )

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer


def create_zip_of_word_docs(submissions_df):
    zip_buffer = io.BytesIO()

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for _, row in submissions_df.iterrows():
            submission = row.to_dict()
            docx_buffer = create_submission_docx(submission)

            filename = (
                clean_filename(submission.get("assignment_title"))
                + "_"
                + clean_filename(submission.get("student_id"))
                + "_"
                + clean_filename(submission.get("student_name"))
                + ".docx"
            )

            zip_file.writestr(filename, docx_buffer.getvalue())

    zip_buffer.seek(0)

    return zip_buffer


# ============================================================
# Teacher login
# ============================================================

def teacher_login(page_key="workflow_teacher"):
    """Compatibility wrapper around the shared teacher-access gate."""
    return require_teacher_access(page_key)



# ============================================================
# Teacher assignment page
# ============================================================

def teacher_assignment_page():
    st.title("Teacher Assignment Creator")

    if not teacher_login("teacher_assignments"):
        st.info("Enter the teacher password to create assignments.")
        return

    st.divider()
    st.subheader("Lecturer Identity")
    st.caption(
        "Assignments store the creator's identity so colleagues can see who created each exercise."
    )
    ident_col1, ident_col2 = st.columns(2)
    with ident_col1:
        creator_name = st.text_input(
            "Lecturer / creator name",
            value=st.session_state.get("eduapp_creator_name", ""),
            placeholder="Example: Dr Nour Abdelaal",
        )
    with ident_col2:
        creator_email = st.text_input(
            "Lecturer email or staff ID",
            value=st.session_state.get("eduapp_creator_email", ""),
            placeholder="Example: nour@university.edu",
        )
    st.session_state["eduapp_creator_name"] = creator_name.strip()
    st.session_state["eduapp_creator_email"] = creator_email.strip()

    st.divider()
    st.subheader("Create a New Assignment")

    with st.form("create_assignment_form"):
        course = st.text_input(
            "Course name",
            placeholder="Example: TRS430",
        )

        title = st.text_input(
            "Assignment title",
            placeholder="Example: Post-editing Task 1",
        )

        st.markdown("**Assigned classes / groups**")
        st.caption(
            "Add one row for each class or group. You may type your own access code or leave "
            "the code blank and EduApp will generate one automatically. Leave the whole table "
            "blank only if the exercise should be visible to all students."
        )
        audience_editor = st.data_editor(
            pd.DataFrame(columns=["Class / group", "Access code"]),
            num_rows="dynamic",
            hide_index=True,
            use_container_width=True,
            key="create_assignment_audience_editor",
            column_config={
                "Class / group": st.column_config.TextColumn(
                    "Class / group",
                    help="Example: TRS430-A",
                ),
                "Access code": st.column_config.TextColumn(
                    "Access code",
                    help="Optional. Leave blank to generate a code automatically.",
                ),
            },
        )

        instructions = st.text_area(
            "Instructions for students",
            placeholder="Explain what students should do.",
            height=120,
        )

        source_text = st.text_area(
            "Source text",
            placeholder="Paste the original text here.",
            height=180,
        )

        machine_translation = st.text_area(
            "Raw machine translation (optional for translation-only assignments)",
            placeholder=(
                "Paste the machine-translated text here. Leave blank when students "
                "should only translate from the source."
            ),
            height=180,
        )

        reference_translation = st.text_area(
            "Reference translation / model answer",
            placeholder="Optional but recommended for quality assessment.",
            height=180,
        )

        due_date = st.date_input("Due date")

        max_score = st.number_input(
            "Maximum score",
            min_value=1.0,
            max_value=100.0,
            value=10.0,
            step=0.5,
        )

        active = st.checkbox(
            "Make this assignment visible to eligible students",
            value=True,
        )

        submitted = st.form_submit_button("Create Assignment")

        if submitted:
            audience_rules, audience_errors = audience_rules_from_editor(audience_editor)
            if not creator_name.strip():
                st.error("Please enter the lecturer / creator name above.")
            elif not title.strip():
                st.error("Please enter an assignment title.")
            elif not source_text.strip():
                st.error("Please enter the source text.")
            elif audience_errors:
                st.error("Please check the class/group table:")
                for error in audience_errors:
                    st.write(f"- {error}")
            else:
                assignment = {
                    "course": course.strip(),
                    "title": title.strip(),
                    "instructions": instructions.strip(),
                    "source_text": source_text.strip(),
                    "machine_translation": machine_translation.strip(),
                    "reference_translation": reference_translation.strip(),
                    "due_date": str(due_date),
                    "max_score": float(max_score),
                    "active": bool(active),
                    "created_by_name": creator_name.strip(),
                    "created_by_email": creator_email.strip(),
                    "audience_rules": audience_rules,
                }

                save_assignment(assignment)
                st.session_state["last_created_assignment_codes"] = {
                    "title": title.strip(),
                    "rules": audience_rules,
                }
                st.success("Assignment created successfully.")
                st.rerun()

    last_created = st.session_state.pop("last_created_assignment_codes", None)
    if last_created:
        st.success(f"Created: {last_created.get('title', 'Assignment')}")
        rules = last_created.get("rules") or []
        if rules:
            st.markdown("### Codes to give your students")
            st.caption("Give each class only its own code.")
            st.dataframe(
                pd.DataFrame(
                    [{"Class / group": r.get("group", ""), "Access code": r.get("code", "")} for r in rules]
                ),
                hide_index=True,
                use_container_width=True,
            )
        else:
            st.info("This assignment is visible to all students; no access code is required.")

    st.divider()
    st.subheader("Existing Assignments")

    assignments = load_assignments()

    if assignments.empty:
        st.info("No assignments created yet.")
        return

    assignments = assignments.copy()
    if "machine_translation" in assignments.columns:
        assignments["student_modes"] = assignments["machine_translation"].apply(
            lambda value: "Translation or post-editing" if safe_text(value) else "Translation only"
        )

    if "audience_rules" in assignments.columns:
        assignments["assigned_groups"] = assignments["audience_rules"].apply(
            lambda value: ", ".join(rule["group"] for rule in parse_audience_rules(value))
            or "All students"
        )
    else:
        assignments["assigned_groups"] = "All students"

    display_columns = [
        "created_at",
        "created_by_name",
        "created_by_email",
        "course",
        "title",
        "assigned_groups",
        "student_modes",
        "due_date",
        "max_score",
        "active",
    ]
    available_columns = [column for column in display_columns if column in assignments.columns]
    st.dataframe(assignments[available_columns], use_container_width=True, hide_index=True)

    st.markdown("### View Access Codes and Edit an Assignment")
    st.caption(
        "Select an exercise to see the exact class/group codes you can give students, "
        "or update the exercise after creation."
    )

    edit_labels = {}
    for _, row in assignments.iterrows():
        assignment_id = safe_text(row.get("assignment_id"))
        edit_label = (
            f"{safe_text(row.get('title')) or 'Untitled'} — "
            f"{safe_text(row.get('course')) or 'No course'} — "
            f"created by {safe_text(row.get('created_by_name')) or 'Legacy/unknown'} — "
            f"ID {assignment_id}"
        )
        edit_labels[edit_label] = assignment_id

    selected_edit_label = st.selectbox(
        "Choose assignment to view or edit",
        list(edit_labels.keys()),
        key="edit_assignment_choice",
    )
    selected_assignment_id = edit_labels[selected_edit_label]
    selected_row = assignments[
        assignments["assignment_id"].astype(str) == str(selected_assignment_id)
    ].iloc[0]

    selected_rules = parse_audience_rules(selected_row.get("audience_rules"))
    if selected_rules:
        st.markdown("#### Student access codes")
        for rule in selected_rules:
            st.code(f"{safe_text(rule.get('group'))} | {safe_text(rule.get('code'))}")
        st.info("Give each class only its own access code. Students use that code on the Student Assignments page.")
    else:
        st.info("This assignment is currently visible to all students and does not require an access code.")

    with st.expander("Edit selected assignment", expanded=False):
        current_rules_text = "\n".join(
            f"{safe_text(rule.get('group'))} | {safe_text(rule.get('code'))}"
            for rule in selected_rules
        )
        current_due = pd.to_datetime(selected_row.get("due_date"), errors="coerce")
        if pd.isna(current_due):
            current_due = pd.Timestamp.today()

        with st.form("edit_assignment_form"):
            edit_course = st.text_input("Course name", value=safe_text(selected_row.get("course")))
            edit_title = st.text_input("Assignment title", value=safe_text(selected_row.get("title")))
            st.markdown("**Assigned classes / groups**")
            st.caption(
                "Edit the table directly. Add or delete rows as needed. If you leave an access "
                "code blank, EduApp will generate a new one when you save."
            )
            edit_audience_editor = st.data_editor(
                pd.DataFrame(
                    [
                        {"Class / group": rule.get("group", ""), "Access code": rule.get("code", "")}
                        for rule in selected_rules
                    ],
                    columns=["Class / group", "Access code"],
                ),
                num_rows="dynamic",
                hide_index=True,
                use_container_width=True,
                key=f"edit_assignment_audience_editor_{selected_assignment_id}",
                column_config={
                    "Class / group": st.column_config.TextColumn("Class / group"),
                    "Access code": st.column_config.TextColumn(
                        "Access code",
                        help="Leave blank to generate a new access code automatically.",
                    ),
                },
            )
            edit_instructions = st.text_area(
                "Instructions for students", value=safe_text(selected_row.get("instructions")), height=120
            )
            edit_source = st.text_area(
                "Source text", value=safe_text(selected_row.get("source_text")), height=180
            )
            edit_mt = st.text_area(
                "Raw machine translation", value=safe_text(selected_row.get("machine_translation")), height=180
            )
            edit_reference = st.text_area(
                "Reference translation / model answer",
                value=safe_text(selected_row.get("reference_translation")),
                height=180,
            )
            edit_due = st.date_input("Due date", value=current_due.date())
            edit_max_score = st.number_input(
                "Maximum score",
                min_value=1.0,
                max_value=100.0,
                value=float(selected_row.get("max_score") or 10.0),
                step=0.5,
            )
            edit_active = st.checkbox(
                "Make this assignment visible to eligible students",
                value=str(selected_row.get("active")).strip().lower() in {"true", "1", "yes"},
            )
            save_changes = st.form_submit_button("Save Changes")

            if save_changes:
                edited_rules, edited_errors = audience_rules_from_editor(edit_audience_editor)
                if not edit_title.strip():
                    st.error("Please enter an assignment title.")
                elif not edit_source.strip():
                    st.error("Please enter the source text.")
                elif edited_errors:
                    st.error("Please check the class/group table:")
                    for error in edited_errors:
                        st.write(f"- {error}")
                else:
                    updates = {
                        "course": edit_course.strip(),
                        "title": edit_title.strip(),
                        "instructions": edit_instructions.strip(),
                        "source_text": edit_source.strip(),
                        "machine_translation": edit_mt.strip(),
                        "reference_translation": edit_reference.strip(),
                        "due_date": str(edit_due),
                        "max_score": float(edit_max_score),
                        "active": bool(edit_active),
                        "audience_rules": edited_rules,
                    }
                    result = update_assignment(selected_assignment_id, updates)
                    if result is not None:
                        st.success("Assignment updated successfully.")
                        st.rerun()

    st.markdown("### Delete an Assignment")
    st.warning(
        "Deleting removes the exercise from the assignment list and deletes unfinished drafts. "
        "Already-submitted student records are preserved for grading and research."
    )
    labels = {}
    for _, row in assignments.iterrows():
        assignment_id = safe_text(row.get("assignment_id"))
        label = (
            f"{safe_text(row.get('title')) or 'Untitled'} — "
            f"{safe_text(row.get('course')) or 'No course'} — "
            f"created by {safe_text(row.get('created_by_name')) or 'Legacy/unknown'} — "
            f"ID {assignment_id}"
        )
        labels[label] = assignment_id

    selected_delete_label = st.selectbox(
        "Choose assignment to delete",
        list(labels.keys()),
        key="delete_assignment_choice",
    )
    confirm_delete = st.checkbox(
        "I understand that this removes the exercise from student access.",
        key="delete_assignment_confirm",
    )
    if st.button("Delete selected assignment", type="primary"):
        if not confirm_delete:
            st.error("Tick the confirmation box before deleting.")
        else:
            result = delete_assignment(labels[selected_delete_label])
            if result is not None:
                st.success("Assignment deleted. Existing submissions were preserved.")
                st.rerun()


# ============================================================
# Student assignment page
# ============================================================

def student_assignment_page():
    st.title("Student Assignments")
    st.write("Choose an assignment, then translate from the source or post-edit the MT output.")

    assignments = load_assignments()

    if assignments.empty:
        st.info("No assignments are available yet.")
        return

    if "active" not in assignments.columns:
        st.info("No active assignments are currently available.")
        return

    active_mask = (
        assignments["active"]
        .fillna(False)
        .astype(str)
        .str.strip()
        .str.lower()
        .isin({"true", "1", "yes"})
    )
    active_assignments = assignments[active_mask].copy()

    if active_assignments.empty:
        st.info("No active assignments are currently available.")
        return

    st.markdown("### Class / Group Access")
    student_access_code = st.text_input(
        "Class or group access code",
        type="password",
        placeholder="Enter the code provided by your lecturer",
        help="Only assignments assigned to your class/group will be shown.",
        key="student_assignment_access_code",
    ).strip()

    def _student_can_access(row):
        rules = parse_audience_rules(row.get("audience_rules")) if "audience_rules" in row.index else []
        if not rules:
            # Backward compatibility: assignments created before group controls remain public.
            return True
        if not student_access_code:
            return False
        return any(
            normalize_access_code(rule.get("code")).casefold() == normalize_access_code(student_access_code).casefold()
            for rule in rules
        )

    active_assignments = active_assignments[
        active_assignments.apply(_student_can_access, axis=1)
    ]

    if active_assignments.empty:
        if student_access_code:
            st.warning("No active assignments are assigned to this class/group access code.")
        else:
            st.info("Enter your class/group access code to view assigned exercises.")
        return

    label_to_id = {}
    for _, row in active_assignments.iterrows():
        assignment_id = safe_text(row.get("assignment_id"))
        label = (
            f"{safe_text(row.get('title')) or 'Untitled'} — "
            f"due {safe_text(row.get('due_date')) or 'not set'} — ID {assignment_id}"
        )
        label_to_id[label] = assignment_id

    selected_label = st.selectbox("Choose an assignment", list(label_to_id))
    selected_assignment_id = label_to_id[selected_label]

    selected_rows = active_assignments[
        active_assignments["assignment_id"].astype(str) == str(selected_assignment_id)
    ]
    if selected_rows.empty:
        st.error("The selected assignment could not be loaded.")
        return

    selected_assignment = selected_rows.iloc[0]

    st.subheader(safe_text(selected_assignment.get("title")) or "Untitled assignment")

    if safe_text(selected_assignment.get("course")):
        st.write(f"**Course:** {selected_assignment.get('course')}")

    st.write(f"**Due date:** {safe_text(selected_assignment.get('due_date')) or 'Not set'}")
    st.write(f"**Maximum score:** {selected_assignment.get('max_score')}")
    selected_rules = parse_audience_rules(selected_assignment.get("audience_rules"))
    matched_group = next(
        (
            rule.get("group")
            for rule in selected_rules
            if normalize_access_code(rule.get("code")).casefold() == normalize_access_code(student_access_code).casefold()
        ),
        "",
    )
    if matched_group:
        st.write(f"**Class / group:** {matched_group}")

    if safe_text(selected_assignment.get("instructions")):
        st.markdown("### Instructions")
        st.write(selected_assignment.get("instructions"))

    source_text = safe_text(selected_assignment.get("source_text"))
    raw_mt = safe_text(selected_assignment.get("machine_translation"))
    reference_translation = safe_text(selected_assignment.get("reference_translation"))

    st.markdown("### Student Information")
    student_col1, student_col2 = st.columns(2)
    with student_col1:
        student_id = st.text_input(
            "Student ID",
            placeholder="Example: S123456",
            key=f"student_id_{selected_assignment_id}",
        )
    with student_col2:
        student_name = st.text_input(
            "Student name",
            placeholder="Example: Aisha Ahmed",
            key=f"student_name_{selected_assignment_id}",
        )

    st.markdown("### Choose the Task Type")

    available_task_labels = list(TASK_OPTIONS)
    if not raw_mt:
        available_task_labels = ["Translate from the source text"]
        st.info(
            "This assignment has no machine translation, so Translation is the only available mode."
        )

    selected_task_label = st.radio(
        "Task type",
        available_task_labels,
        horizontal=True,
        key=f"task_type_{selected_assignment_id}",
    )
    task_type = TASK_OPTIONS[selected_task_label]
    st.info(task_instruction(task_type))

    answer_key = (
        f"student_translation_{selected_assignment_id}"
        if is_translation(task_type)
        else f"student_post_edit_{selected_assignment_id}"
    )
    default_answer = "" if is_translation(task_type) else raw_mt
    if answer_key not in st.session_state:
        st.session_state[answer_key] = default_answer

    st.markdown("### Continue a Saved Draft")
    st.caption(
        "Drafts are stored in Supabase and are separate for Translation and Post-editing. "
        "Enter the same Student ID when you return."
    )
    if st.button(
        "Load / resume saved draft",
        key=f"load_draft_{selected_assignment_id}_{task_type}",
    ):
        if not student_id.strip():
            st.error("Enter your Student ID first so the app can find your draft.")
        else:
            draft = load_student_draft(selected_assignment_id, task_type, student_id)
            if draft:
                st.session_state[answer_key] = safe_text(draft.get("draft_text"))
                st.session_state[f"draft_loaded_notice_{selected_assignment_id}_{task_type}"] = True
                st.rerun()
            else:
                st.info("No saved draft was found for this Student ID and task type.")

    if st.session_state.pop(
        f"draft_loaded_notice_{selected_assignment_id}_{task_type}", False
    ):
        st.success("Saved draft loaded. You can continue from where you stopped.")

    # Keep the source and the student's working box together so students can compare
    # them without scrolling up and down.
    source_col, work_col = st.columns(2, gap="medium")
    with source_col:
        st.markdown("### Source Text")
        st.text_area(
            "Source text",
            source_text,
            height=360,
            disabled=True,
            label_visibility="collapsed",
        )

    if is_translation(task_type):
        with work_col:
            st.markdown("### Your Translation")
            student_answer = st.text_area(
                "Translation box",
                key=answer_key,
                height=360,
                placeholder="Write your translation here.",
                label_visibility="collapsed",
            )
        edit_summary = {
            "inserted_words": None,
            "deleted_words": None,
            "replaced_segments": None,
            "unchanged_words": None,
        }
    else:
        with work_col:
            st.markdown("### Post-edit the MT Output")
            st.caption("Raw MT is shown below for reference while you post-edit.")
            st.text_area(
                "Original raw MT output",
                raw_mt,
                height=150,
                disabled=True,
                label_visibility="collapsed",
            )
            student_answer = st.text_area(
                "Post-editing box",
                key=answer_key,
                height=170,
                label_visibility="collapsed",
            )

        st.markdown("### Track Changes Preview")
        track_changes_html = make_track_changes_html(raw_mt, student_answer)
        st.markdown(
            f'<div class="track-box">{track_changes_html}</div>',
            unsafe_allow_html=True,
        )

        edit_summary = calculate_edit_summary(raw_mt, student_answer)
        st.markdown("### Editing Summary")
        st.dataframe(
            pd.DataFrame(
                [
                    {"Edit feature": "Inserted words", "Value": edit_summary["inserted_words"]},
                    {"Edit feature": "Deleted words", "Value": edit_summary["deleted_words"]},
                    {"Edit feature": "Replaced segments", "Value": edit_summary["replaced_segments"]},
                    {"Edit feature": "Unchanged words", "Value": edit_summary["unchanged_words"]},
                ]
            ),
            use_container_width=True,
            hide_index=True,
        )

    save_col, note_col = st.columns([1, 2])
    with save_col:
        save_draft_clicked = st.button(
            "Save draft for later",
            key=f"save_draft_{selected_assignment_id}_{task_type}",
            use_container_width=True,
        )
    with note_col:
        st.caption(
            "Saving a draft does not submit the assignment. You may close the app and return later."
        )

    if save_draft_clicked:
        if not student_id.strip():
            st.error("Enter your Student ID before saving a draft.")
        else:
            result = save_student_draft(
                {
                    "assignment_id": safe_text(selected_assignment_id),
                    "task_type": normalize_task_type(task_type),
                    "student_id": student_id.strip(),
                    "student_name": student_name.strip(),
                    "draft_text": student_answer,
                }
            )
            if result is not None:
                st.success(
                    "Draft saved. Return to this assignment later, enter the same Student ID, "
                    "choose the same task type, and click 'Load / resume saved draft'."
                )

    # Apply the browser-side integrity guard only to the assessed student response boxes.
    install_student_paste_guard()

    with st.expander("Advanced metric settings (research or pilot use)", expanded=False):
        research_mode = st.toggle(
            "Research mode",
            value=True,
            help="Stores research fields and marks advanced neural metrics as pending.",
            key=f"research_mode_{selected_assignment_id}_{task_type}",
        )
        run_advanced_now = st.toggle(
            "Run advanced metrics now",
            value=False,
            help=(
                "Leave this off during live classes. BERTScore and COMET can be "
                "calculated later in batch."
            ),
            key=f"advanced_now_{selected_assignment_id}_{task_type}",
        )
        use_semantic_cosine = st.checkbox(
            "Use semantic cosine similarity",
            value=False,
            help="Semantic cosine may be slower the first time the model loads.",
            key=f"semantic_cosine_{selected_assignment_id}_{task_type}",
        )
        use_bert = st.checkbox(
            "Calculate BERTScore",
            value=False,
            key=f"bert_{selected_assignment_id}_{task_type}",
        )
        bert_language = st.selectbox(
            "BERTScore language",
            ["en", "ar", "fr", "de", "es", "zh", "ja", "ko", "tr", "ru"],
            index=0,
            key=f"bert_language_{selected_assignment_id}_{task_type}",
        )

    reference_metrics = compare_postedit_and_reference(
        post_edited_text=student_answer,
        reference_translation=reference_translation,
        use_semantic_cosine=use_semantic_cosine,
        use_bert=False,
        bert_language=bert_language,
    )

    if is_translation(task_type):
        mt_pe_metrics = {
            "mt_pe_cosine_similarity": None,
            "mt_pe_cosine_method": "not_applicable_translation_task",
            "mt_pe_edit_distance_ratio": None,
            "mt_pe_length_ratio": None,
            "mt_pe_bleu": None,
            "mt_pe_chrf": None,
            "mt_pe_ter": None,
            "mt_pe_bertscore_f1": None,
        }
        if reference_translation:
            st.markdown("### Translation vs Reference Translation")
            st.dataframe(
                metrics_to_dataframe(reference_metrics),
                use_container_width=True,
                hide_index=True,
            )
        else:
            st.warning(
                "No reference translation was provided, so automatic quality metrics are limited."
            )
    else:
        mt_pe_metrics = compare_mt_and_postedit(
            raw_mt=raw_mt,
            post_edited_text=student_answer,
            use_semantic_cosine=use_semantic_cosine,
            use_bert=False,
            bert_language=bert_language,
        )
        st.markdown("### Raw MT vs Post-Edited Text")
        st.dataframe(
            metrics_to_dataframe(mt_pe_metrics),
            use_container_width=True,
            hide_index=True,
        )
        st.caption(
            "MT–PE metrics indicate editing overlap and effort; they are not final quality scores."
        )

        if reference_translation:
            st.markdown("### Post-Edited Text vs Reference Translation")
            st.dataframe(
                metrics_to_dataframe(reference_metrics),
                use_container_width=True,
                hide_index=True,
            )
        else:
            st.warning(
                "No reference translation was provided, so reference-based quality metrics are limited."
            )

    button_label = f"Submit {task_type_label(task_type)} Task"
    if st.button(button_label, type="primary", key=f"submit_{selected_assignment_id}_{task_type}"):
        if not student_id.strip():
            st.error("Please enter your student ID.")
            return

        if not student_answer.strip():
            st.error(f"Please enter your {student_output_label(task_type).lower()} before submitting.")
            return

        with st.spinner("Calculating final metrics and saving submission..."):
            if not is_translation(task_type):
                mt_pe_metrics = compare_mt_and_postedit(
                    raw_mt=raw_mt,
                    post_edited_text=student_answer,
                    use_semantic_cosine=use_semantic_cosine,
                    use_bert=use_bert,
                    bert_language=bert_language,
                )

            reference_metrics = compare_postedit_and_reference(
                post_edited_text=student_answer,
                reference_translation=reference_translation,
                use_semantic_cosine=use_semantic_cosine,
                use_bert=use_bert,
                bert_language=bert_language,
            )

            output_word_count = word_count(student_answer)
            quality_warnings = build_quality_warnings(
                mt_pe_metrics,
                reference_metrics,
                output_word_count,
                task_type=task_type,
            )

            research_results = compare_postedit_with_raw_mt(
                raw_mt=raw_mt,
                post_edited_text=student_answer,
                human_translation=student_answer if is_translation(task_type) else None,
                reference_text=reference_translation,
                source_text=source_text,
                teacher_score=None,
                teacher_feedback="",
                use_bert=run_advanced_now and use_bert,
                bert_language=bert_language,
                comet_scorer=None,
            )

            if is_translation(task_type):
                make_translation_metrics_task_appropriate(
                    research_results,
                    has_reference=bool(reference_translation),
                )

            submission = {
                "assignment_id": selected_assignment.get("assignment_id"),
                "assignment_title": selected_assignment.get("title"),
                "task_type": normalize_task_type(task_type),
                "student_id": student_id.strip(),
                "student_name": student_name.strip(),
                "group_name": matched_group,
                "source_text": source_text,
                # The MT may remain stored for teacher/research comparison, but it is hidden
                # from students in Translation mode.
                "machine_translation": raw_mt,
                "reference_translation": reference_translation,
                # Keep this legacy column as the canonical final student output so that
                # the existing annotation, AI, and analytics pages remain compatible.
                "post_edited_text": student_answer.strip(),
                "inserted_words": edit_summary["inserted_words"],
                "deleted_words": edit_summary["deleted_words"],
                "replaced_segments": edit_summary["replaced_segments"],
                "unchanged_words": edit_summary["unchanged_words"],
                "source_word_count": word_count(source_text),
                "mt_word_count": word_count(raw_mt) if not is_translation(task_type) else None,
                "pe_word_count": output_word_count,
                "quality_warnings": quality_warnings,
                "teacher_score": None,
                "teacher_feedback": "",
            }

            submission.update(mt_pe_metrics)
            submission.update(reference_metrics)
            submission.update(
                build_research_metrics_payload(
                    research_results,
                    research_mode=research_mode,
                )
            )

            save_submission(submission)
            delete_student_draft(selected_assignment_id, task_type, student_id)

        st.success(f"Your {task_type_label(task_type).lower()} submission has been saved.")

        st.subheader("Automatic Assessment Indicators")
        if reference_translation:
            st.dataframe(
                metrics_to_dataframe(reference_metrics),
                use_container_width=True,
                hide_index=True,
            )
        st.write(f"**Automatic warnings:** {quality_warnings}")
        st.warning(
            "Automatic metrics are indicators only. The teacher makes the final assessment."
        )


# ============================================================
# Teacher submissions dashboard
# ============================================================

def teacher_submissions_page():
    st.title("Teacher Submissions Dashboard")

    if not teacher_login("teacher_submissions"):
        st.info("Enter the teacher password to view submissions.")
        return

    submissions = load_submissions()

    if submissions.empty:
        st.info("No student submissions yet.")
        return

    assignments = load_assignments()

    assignment_titles = sorted(
        submissions["assignment_title"].dropna().astype(str).unique().tolist()
    )

    selected_assignment_title = st.selectbox(
        "Choose assignment",
        assignment_titles,
    )

    filtered = submissions[
        submissions["assignment_title"].astype(str) == selected_assignment_title
    ].copy()

    if "task_type" not in filtered.columns:
        filtered["task_type"] = POST_EDITING
    filtered["task_type"] = filtered["task_type"].apply(normalize_task_type)

    task_filter = st.selectbox(
        "Filter by task type",
        ["All task types", "Translation", "Post-editing"],
    )
    if task_filter != "All task types":
        wanted = TRANSLATION if task_filter == "Translation" else POST_EDITING
        filtered = filtered[filtered["task_type"] == wanted]

    if filtered.empty:
        st.info("No submissions match this assignment and task-type filter.")
        return

    st.subheader(f"Submissions for: {selected_assignment_title}")

    display_columns = [
        "submitted_at",
        "student_id",
        "student_name",
        "task_type",
        "pe_word_count",
        "inserted_words",
        "deleted_words",
        "replaced_segments",
        "mt_pe_cosine_similarity",
        "mt_pe_edit_distance_ratio",
        "mt_pe_chrf",
        "mt_pe_ter",
        "pe_reference_cosine_similarity",
        "pe_reference_chrf",
        "pe_reference_ter",
        "research_mode",
        "advanced_metrics_status",
        "pe_quality_chrf",
        "pe_quality_ter",
        "quality_warnings",
        "teacher_score",
    ]

    available_columns = [
        column for column in display_columns if column in filtered.columns
    ]

    st.dataframe(
        filtered[available_columns],
        use_container_width=True,
        hide_index=True,
    )

    st.divider()

    st.subheader("Research Export")

    research_columns = [
        "submitted_at",
        "assignment_id",
        "assignment_title",
        "task_type",
        "student_id",
        "student_name",
        "source_text",
        "machine_translation",
        "reference_translation",
        "post_edited_text",
        "inserted_words",
        "deleted_words",
        "replaced_segments",
        "unchanged_words",
        "source_word_count",
        "mt_word_count",
        "pe_word_count",
        "mt_pe_cosine_similarity",
        "mt_pe_cosine_method",
        "mt_pe_edit_distance_ratio",
        "mt_pe_length_ratio",
        "mt_pe_bleu",
        "mt_pe_chrf",
        "mt_pe_ter",
        "mt_pe_bertscore_f1",
        "pe_reference_cosine_similarity",
        "pe_reference_cosine_method",
        "pe_reference_length_ratio",
        "pe_reference_bleu",
        "pe_reference_chrf",
        "pe_reference_ter",
        "pe_reference_bertscore_f1",
        "research_mode",
        "advanced_metrics_status",
        "raw_mt_word_count",
        "reference_word_count",
        "mt_pe_word_count_difference",
        "mt_pe_lexical_similarity",
        "mt_pe_change_ratio",
        "mt_pe_replacement_output_words",
        "mt_pe_changed_original_words",
        "mt_pe_unchanged_ratio",
        "mt_pe_changed_ratio_original",
        "mt_pe_overlap_bleu",
        "mt_pe_overlap_chrf",
        "mt_pe_overlap_ter",
        "raw_mt_quality_bleu",
        "raw_mt_quality_chrf",
        "raw_mt_quality_ter",
        "pe_quality_bleu",
        "pe_quality_chrf",
        "pe_quality_ter",
        "ht_quality_bleu",
        "ht_quality_chrf",
        "ht_quality_ter",
        "raw_mt_quality_bertscore_f1",
        "pe_quality_bertscore_f1",
        "ht_quality_bertscore_f1",
        "raw_mt_quality_comet",
        "pe_quality_comet",
        "ht_quality_comet",
        "mt_pe_interpretation",
        "quality_warnings",
        "teacher_score",
        "teacher_feedback",
    ]

    available_research_columns = [
        column for column in research_columns if column in filtered.columns
    ]

    research_df = filtered[available_research_columns]

    csv_data = research_df.to_csv(index=False).encode("utf-8")

    st.download_button(
        "Download research dataset as CSV",
        data=csv_data,
        file_name="translation_postediting_research_dataset.csv",
        mime="text/csv",
    )

    excel_buffer = io.BytesIO()

    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        research_df.to_excel(writer, index=False, sheet_name="Research Data")

    excel_buffer.seek(0)

    st.download_button(
        "Download research dataset as Excel",
        data=excel_buffer,
        file_name="translation_postediting_research_dataset.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    zip_buffer = create_zip_of_word_docs(filtered)

    st.download_button(
        "Download all submissions as Word documents",
        data=zip_buffer,
        file_name="translation_postediting_submissions.zip",
        mime="application/zip",
    )

    st.divider()

    st.subheader("Review Individual Submission")

    submission_labels = []

    for _, row in filtered.iterrows():
        label = (
            f"{row.get('student_id', '')} — {row.get('student_name', '')} — "
            f"{task_type_label(row.get('task_type'))} — "
            f"{row.get('submitted_at', '')} — ID {row.get('submission_id', '')}"
        )
        submission_labels.append(label)

    selected_submission_label = st.selectbox(
        "Choose submission",
        submission_labels,
    )

    selected_submission_id = selected_submission_label.split("ID ")[-1]

    selected_submission = filtered[
        filtered["submission_id"].astype(str) == selected_submission_id
    ].iloc[0]

    selected_task_type = normalize_task_type(selected_submission.get("task_type"))
    st.write(f"**Task type:** {task_type_label(selected_task_type)}")

    st.markdown("### Source Text")
    st.text_area(
        "Selected source text",
        selected_submission.get("source_text", ""),
        height=180,
        disabled=True,
        label_visibility="collapsed",
    )

    if not is_translation(selected_task_type):
        st.markdown("### Raw Machine Translation")
        st.text_area(
            "Selected raw MT",
            selected_submission.get("machine_translation", ""),
            height=180,
            disabled=True,
            label_visibility="collapsed",
        )

    output_label = student_output_label(selected_task_type)
    st.markdown(f"### {output_label}")
    st.text_area(
        output_label,
        selected_submission.get("post_edited_text", ""),
        height=250,
        disabled=True,
        label_visibility="collapsed",
    )

    if not is_translation(selected_task_type):
        st.markdown("### Track Changes Preview")
        st.markdown(
            f"""
            <div class="track-box">
            {make_track_changes_html(
                selected_submission.get("machine_translation", ""),
                selected_submission.get("post_edited_text", ""),
            )}
            </div>
            """,
            unsafe_allow_html=True,
        )

    metric_columns = [
        "inserted_words",
        "deleted_words",
        "replaced_segments",
        "unchanged_words",
        "source_word_count",
        "mt_word_count",
        "pe_word_count",
        "mt_pe_cosine_similarity",
        "mt_pe_cosine_method",
        "mt_pe_edit_distance_ratio",
        "mt_pe_length_ratio",
        "mt_pe_bleu",
        "mt_pe_chrf",
        "mt_pe_ter",
        "mt_pe_bertscore_f1",
        "pe_reference_cosine_similarity",
        "pe_reference_cosine_method",
        "pe_reference_length_ratio",
        "pe_reference_bleu",
        "pe_reference_chrf",
        "pe_reference_ter",
        "pe_reference_bertscore_f1",
        "research_mode",
        "advanced_metrics_status",
        "raw_mt_quality_bleu",
        "raw_mt_quality_chrf",
        "raw_mt_quality_ter",
        "pe_quality_bleu",
        "pe_quality_chrf",
        "pe_quality_ter",
        "ht_quality_bleu",
        "ht_quality_chrf",
        "ht_quality_ter",
        "raw_mt_quality_bertscore_f1",
        "pe_quality_bertscore_f1",
        "ht_quality_bertscore_f1",
        "raw_mt_quality_comet",
        "pe_quality_comet",
        "ht_quality_comet",
        "mt_pe_interpretation",
        "quality_warnings",
    ]

    if is_translation(selected_task_type):
        post_editing_only = set(POST_EDITING_ONLY_METRIC_FIELDS) | {
            "inserted_words",
            "deleted_words",
            "replaced_segments",
            "unchanged_words",
            "mt_word_count",
        }
        metric_columns = [
            column for column in metric_columns if column not in post_editing_only
        ]

    metric_rows = []

    for column in metric_columns:
        metric_rows.append(
            {
                "Metric": column,
                "Value": selected_submission.get(column, ""),
            }
        )

    st.markdown("### Automatic Metrics")

    st.dataframe(
        pd.DataFrame(metric_rows),
        use_container_width=True,
        hide_index=True,
    )

    st.markdown("### Teacher Review")

    max_score = 100.0

    if not assignments.empty and "title" in assignments.columns:
        matching_assignment = assignments[
            assignments["title"].astype(str) == str(selected_assignment_title)
        ]

        if not matching_assignment.empty:
            try:
                max_score = float(matching_assignment.iloc[0]["max_score"])
            except Exception:
                max_score = 100.0

    current_score = selected_submission.get("teacher_score", 0)

    try:
        current_score = float(current_score)
    except Exception:
        current_score = 0.0

    teacher_score = st.number_input(
        "Teacher score",
        min_value=0.0,
        max_value=max_score,
        value=current_score,
        step=0.5,
    )

    teacher_feedback = st.text_area(
        "Teacher feedback",
        value=safe_text(selected_submission.get("teacher_feedback")),
        height=120,
    )

    if st.button("Save Teacher Review"):
        update_submission_review(
            selected_submission_id,
            teacher_score,
            teacher_feedback,
        )

        st.success("Teacher review saved.")
        st.rerun()

    single_docx = create_submission_docx(selected_submission.to_dict())

    single_filename = (
        clean_filename(selected_assignment_title)
        + "_"
        + clean_filename(selected_submission.get("student_id"))
        + ".docx"
    )

    st.download_button(
        "Download this submission as Word document",
        data=single_docx,
        file_name=single_filename,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
